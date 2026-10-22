"""Convert an Excel question file into a question-set file for the website.

Usage:
    python tools/build_set.py files/CorpFin_300.xlsx --id corpfin --title "Корпоративные финансы"
    python tools/build_set.py path/to/Other.xlsx --id macro --title "Макроэкономика" --description "..."

Expected layout (same as CorpFin_300.xlsx):
    Question sheet (named "Лист1" or the first sheet), columns in this order:
        №, Вопрос, Вариант A, Вариант B, Вариант C, Вариант D,
        Правильный вариант (A–D), Категория, Подсказка, Сложность (1–5), Мнемоника
    Optional sheet "Источники": №, Источник, Где используется, Ссылка

The script writes sets/<id>.js and rebuilds sets/manifest.js from every set file
in the folder. Full row validation runs in the browser on load; errors are shown
in the admin panel ("Импорт" tab). The quick checks here only warn.
"""

import argparse
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SETS_DIR = ROOT / "sets"
QUESTION_COLS = 11
SOURCE_COLS = 4


def cell(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        return v.strip()
    return v


def data_rows(ws, width):
    rows = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        r = [cell(v) for v in list(r)[:width]] + [""] * max(0, width - len(r))
        if all(v == "" for v in r):
            continue
        # The header row has no number in the first column.
        if i == 0 and not isinstance(r[0], int):
            continue
        rows.append(r)
    return rows


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("--id", required=True, help="латиница, цифры, дефис: corpfin, macro-2026")
    ap.add_argument("--title", required=True)
    ap.add_argument("--description", default="")
    args = ap.parse_args()

    if not re.fullmatch(r"[a-z0-9][a-z0-9-]{0,40}", args.id):
        sys.exit("--id: только строчная латиница, цифры и дефис")

    wb = openpyxl.load_workbook(args.xlsx, data_only=True, read_only=True)
    qsheet = wb["Лист1"] if "Лист1" in wb.sheetnames else wb.worksheets[0]
    questions = data_rows(qsheet, QUESTION_COLS)
    sources = data_rows(wb["Источники"], SOURCE_COLS) if "Источники" in wb.sheetnames else []

    bad = [r[0] for r in questions if str(r[6]).upper() not in ("A", "B", "C", "D") or "" in r[1:8]]
    if bad:
        print(f"Внимание: строки с неполными данными (будут пропущены при загрузке): {bad[:20]}")

    payload = {
        "id": args.id,
        "title": args.title,
        "description": args.description,
        "sourceFile": Path(args.xlsx).name,
        "questions": questions,
        "sources": sources,
    }
    SETS_DIR.mkdir(exist_ok=True)
    out = SETS_DIR / f"{args.id}.js"
    out.write_text(
        "// Сгенерировано tools/build_set.py — правьте Excel-файл и перезапустите скрипт.\n"
        f"registerQuestionSet({json.dumps(payload, ensure_ascii=False, indent=1)});\n",
        encoding="utf-8",
    )

    files = sorted(p.name for p in SETS_DIR.glob("*.js") if p.name != "manifest.js")
    (SETS_DIR / "manifest.js").write_text(
        "// Список файлов наборов вопросов. Обновляется tools/build_set.py автоматически;\n"
        "// чтобы убрать набор с сайта, удалите его файл и строку здесь.\n"
        f"window.SET_MANIFEST = {json.dumps(files, ensure_ascii=False, indent=2)};\n",
        encoding="utf-8",
    )
    print(f"{out.relative_to(ROOT)}: {len(questions)} вопросов, {len(sources)} источников")
    print(f"sets/manifest.js: {files}")


if __name__ == "__main__":
    main()
